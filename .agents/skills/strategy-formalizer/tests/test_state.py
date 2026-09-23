"""
Agents: read the English part only. 中文仅供人类阅读。

State-machine tests: init, start/complete/accept ordering, user-edit
detection via checkpoint hashes, reopen resetting downstream stages.

状态机测试：init、start/complete/accept 的顺序约束、通过检查点哈希检测用户编辑、reopen
重置下游阶段。
"""
from __future__ import annotations

from pathlib import Path

import pytest

import sf_state
from conftest import load


def _init(tmp_path: Path, monkeypatch) -> Path:
    """
    Create a fake input folder and initialise a workspace under tmp_path, redirecting ACTIVE so the real workspace/ is untouched.

    创建假输入目录并在 tmp_path 下初始化 workspace，重定向 ACTIVE 以免触碰真实的 workspace/。
    """
    inp = tmp_path / "input"
    inp.mkdir()
    (inp / "raw.md").write_text("策略原文", encoding="utf-8")
    ws = tmp_path / "ws"
    monkeypatch.setattr(sf_state, "ACTIVE_FILE", tmp_path / "ACTIVE")
    monkeypatch.setattr(sf_state, "WORKSPACE_ROOT", tmp_path)
    sf_state.main(["--workspace", str(ws), "init", "--input", str(inp), "--name", "demo"])
    return ws


def test_init_and_order(tmp_path, monkeypatch, capsys):
    """
    init creates state.json with S1 current; S2 cannot start before S1 is done.

    init 生成以 S1 为当前阶段的 state.json；S1 未 done 时 S2 不能开始。
    """
    ws = _init(tmp_path, monkeypatch)
    st = load(ws / "state.json")
    assert st["current_stage"] == "S1" and st["stages"]["S1"]["status"] == "pending"
    with pytest.raises(SystemExit):
        sf_state.main(["--workspace", str(ws), "start", "S2"])
    sf_state.main(["--workspace", str(ws), "start", "S1"])
    with pytest.raises(SystemExit):  # outputs missing
        sf_state.main(["--workspace", str(ws), "complete", "S1"])


def test_complete_accept_detects_edits(tmp_path, monkeypatch, capsys):
    """
    complete records hashes; a user edit during the stop is reported by
    status and accept; accept advances to S2.

    complete 记录哈希；停顿期间的用户编辑会被 status 与 accept 报告；accept 推进到 S2。
    """
    ws = _init(tmp_path, monkeypatch)
    sf_state.main(["--workspace", str(ws), "start", "S1"])
    for f in ("S1_sources.json", "S1_strategy_raw.md", "S1_corrections.md", "S1_strategy_clean.md"):
        (ws / f).write_text("{}" if f.endswith("json") else "x", encoding="utf-8")
    sf_state.main(["--workspace", str(ws), "complete", "S1"])
    assert load(ws / "state.json")["stages"]["S1"]["status"] == "awaiting_review"
    (ws / "S1_strategy_raw.md").write_text("edited by user", encoding="utf-8")
    capsys.readouterr()
    sf_state.main(["--workspace", str(ws), "status"])
    assert "S1_strategy_raw.md" in capsys.readouterr().out
    sf_state.main(["--workspace", str(ws), "accept", "S1"])
    st = load(ws / "state.json")
    assert st["stages"]["S1"]["status"] == "done" and st["current_stage"] == "S2"


def test_reopen_resets_downstream(tmp_path, monkeypatch):
    """
    reopen S1 after S1 done puts S1 in_progress and requires --yes.

    S1 完成后 reopen S1 需要 --yes，并把 S1 置回 in_progress。
    """
    ws = _init(tmp_path, monkeypatch)
    sf_state.main(["--workspace", str(ws), "start", "S1"])
    for f in ("S1_sources.json", "S1_strategy_raw.md", "S1_corrections.md", "S1_strategy_clean.md"):
        (ws / f).write_text("x", encoding="utf-8")
    sf_state.main(["--workspace", str(ws), "complete", "S1"])
    sf_state.main(["--workspace", str(ws), "accept", "S1"])
    with pytest.raises(SystemExit):
        sf_state.main(["--workspace", str(ws), "reopen", "S1"])
    sf_state.main(["--workspace", str(ws), "reopen", "S1", "--yes"])
    st = load(ws / "state.json")
    assert st["stages"]["S1"]["status"] == "in_progress" and st["stages"]["S2"]["status"] == "pending"


def test_xlsx_resave_not_flagged(tmp_path, monkeypatch):
    """
    Re-saving an xlsx with identical cells (what Excel does on open+save)
    does not count as a user edit; changing a cell does.

    以相同单元格重新保存 xlsx（Excel 打开再保存的效果）不算用户修改；改动单元格才算。
    """
    import json
    from openpyxl import Workbook, load_workbook
    ws = _init(tmp_path, monkeypatch)
    wb = Workbook(); wb.active.append(["a", None, 1]); wb.save(ws / "S1_extra.xlsx")
    st = load(ws / "state.json")
    st["stages"]["S1"]["checkpoint_hashes"]["S1_extra.xlsx"] = sf_state.content_hash(ws / "S1_extra.xlsx")
    (ws / "state.json").write_text(json.dumps(st), encoding="utf-8")
    load_workbook(ws / "S1_extra.xlsx").save(ws / "S1_extra.xlsx")
    assert "S1_extra.xlsx" not in sf_state.modified_since_checkpoint(ws, load(ws / "state.json"), "S1")
    wb2 = load_workbook(ws / "S1_extra.xlsx"); wb2.active["A1"] = "b"; wb2.save(ws / "S1_extra.xlsx")
    assert "S1_extra.xlsx" in sf_state.modified_since_checkpoint(ws, load(ws / "state.json"), "S1")


def test_complete_last_stage_is_done(tmp_path, monkeypatch):
    """
    complete S8 marks the stage done directly (no accept), because
    freezing final/ ends the run.

    complete S8 直接标记为 done（不需要 accept），因为冻结 final/ 即结束。
    """
    ws = _init(tmp_path, monkeypatch)
    st = load(ws / "state.json")
    for s_ in sf_state.STAGES[:-1]:
        st["stages"][s_]["status"] = "done"
    st["current_stage"] = "S8"
    (ws / "state.json").write_text(__import__("json").dumps(st), encoding="utf-8")
    sf_state.main(["--workspace", str(ws), "start", "S8"])
    (ws / "final").mkdir()
    for f in ("S8_check_report.md", "final/strategy.md", "final/terms.json", "final/formal.json", "final/callbacks_stub.py"):
        (ws / f).write_text("x", encoding="utf-8")
    sf_state.main(["--workspace", str(ws), "complete", "S8"])
    assert load(ws / "state.json")["stages"]["S8"]["status"] == "done"
