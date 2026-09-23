"""
Agents: read the English part only. 中文仅供人类阅读。

Term-file tests: validation of the fixture, xlsx round-trip with a user
edit and a new row, and error detection (hard limit, dangling supports).

术语文件测试：夹具校验、带用户编辑和新行的 xlsx 往返、错误检测（硬上限、悬空 supports）。
"""
from __future__ import annotations

import copy

import pytest
from openpyxl import load_workbook

import sf_terms
from conftest import load


def test_fixture_valid(ws):
    """
    The fixture term file passes the strictest validation (classified + defined).

    夹具术语文件通过最严格的校验（已分类 + 已定义）。
    """
    data = sf_terms.load_terms(ws / "S6_terms_defined.json")
    errors, warnings = sf_terms.validate(data, None, True, True)
    assert errors == []


def test_xlsx_roundtrip(ws):
    """
    Editing a cell and appending an id-less row in the xlsx is merged back: the edit is recorded in history and the row gets the next id with origin user_edit.

    在 xlsx 中修改单元格并追加无 id 的行后能合并回来：修改记入 history，新行获得下一个 id 且 origin 为 user_edit。
    """
    w = str(ws)
    sf_terms.main(["--workspace", w, "to-xlsx", "S6_terms_defined.json"])
    xlsx = ws / "S6_terms_defined.xlsx"
    wb = load_workbook(xlsx)
    sh = wb["terms"]
    header = [c.value for c in sh[1]]
    col = header.index("definition_zh") + 1
    sh.cell(row=2, column=col, value="用户修改后的定义")
    # new row without id
    row = [""] * len(header)
    row[header.index("name_zh")] = "初始止损位置"
    row[header.index("name_en")] = "initial_stop_level"
    row[header.index("role")] = "aux_note"
    row[header.index("phase")] = "position_mgmt"
    row[header.index("supports")] = "T004"
    row[header.index("status")] = "kept"
    row[header.index("definition_zh")] = "小周期结构下方"
    sh.append(row)
    wb.save(xlsx)
    sf_terms.main(["--workspace", w, "from-xlsx", "S6_terms_defined.xlsx"])
    data = load(ws / "S6_terms_defined.json")
    t1 = data["terms"][0]
    assert t1["definition_zh"] == "用户修改后的定义"
    assert any(h["field"] == "definition_zh" and h["by"] == "user_edit" for h in t1["history"])
    new = data["terms"][-1]
    assert new["id"] == "T013" and new["origin"] == "user_edit" and new["supports"] == ["T004"]


def test_validate_errors(ws):
    """
    Dangling and dropped supports targets and more than 50 live terms are reported as errors.

    supports 指向不存在或已丢弃的术语、以及超过 50 个存活术语都会报错。
    """
    data = sf_terms.load_terms(ws / "S6_terms_defined.json")
    bad = copy.deepcopy(data)
    bad["terms"][6]["supports"] = ["T999"]          # dangling
    bad["terms"][7]["supports"] = ["T012"]          # dropped target
    errors, _ = sf_terms.validate(bad)
    assert any("unknown id T999" in e for e in errors)
    assert any("dropped term T012" in e for e in errors)
    many = copy.deepcopy(data)
    for i in range(60):
        t = copy.deepcopy(data["terms"][6])
        t["id"] = f"T{100 + i}"
        t["name_zh"] = f"术语{i}"
        t["name_en"] = f"term_{i}"
        many["terms"].append(t)
    errors, _ = sf_terms.validate(many)
    assert any("hard limit" in e for e in errors)


def test_carry_and_diff(ws, capsys):
    """
    carry copies a term file forward, add appends a term, diff shows it as added.

    carry 向前复制术语文件，add 追加术语，diff 将其显示为新增。
    """
    w = str(ws)
    sf_terms.main(["--workspace", w, "carry", "S6_terms_defined.json", "final/terms.json", "--stage", "S8"])
    assert (ws / "final" / "terms.json").exists()
    sf_terms.main(["--workspace", w, "add", "final/terms.json", "--name-zh", "新术语", "--name-en", "new_term",
                   "--role", "aux_note", "--phase", "entry", "--supports", "T002"])
    capsys.readouterr()
    sf_terms.main(["--workspace", w, "diff", "S6_terms_defined.json", "final/terms.json"])
    assert "+ T013 新术语" in capsys.readouterr().out
