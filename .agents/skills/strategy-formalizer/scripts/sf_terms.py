#!/usr/bin/env python3
"""
Agents: read the English part of every docstring only. 中文段落仅供人类阅读。

sf_terms.py — create, convert, validate and diff term files.

A term file (S2_terms_init.json, S4_terms_filtered.json, S5_terms_classified.json,
S6_terms_defined.json, final/terms.json) is the same structure carried forward
stage by stage: {"stage", "scheme_id", "terms": [term, ...]}. A term has a
stable id T###, Chinese and English names, a status (candidate/kept/dropped/
defined), a role and a phase (assigned in S5), the source quote it was
extracted from, definitions, the ids of key terms it `supports` (for
auxiliary terms), the summary sections it appears in, its origin, a drop
reason, free notes, and a json-only history[] of changes. Terms are never
deleted; they are marked dropped so S4/S6 decisions stay traceable.

The .xlsx twin is the human edit surface (one row per term, list cells joined
by "; ", a legend sheet explaining columns and category values). `to-xlsx`
renders it; `from-xlsx` merges edits back into the json, appending history
entries and assigning ids to rows that have none. json is authoritative.

    sf_terms.py new       S2_terms_init.json --stage S2
    sf_terms.py add       S6_terms_defined.json --name-zh 有利运动 --name-en favorable_move --origin S6_dialog [...]
    sf_terms.py carry     S2_terms_init.json S4_terms_filtered.json --stage S4
    sf_terms.py to-xlsx   S4_terms_filtered.json [--categories S5_categories.json]
    sf_terms.py from-xlsx S4_terms_filtered.xlsx
    sf_terms.py validate  S5_terms_classified.json [--categories S5_categories.json] [--require-classified] [--require-defined]
    sf_terms.py diff      S2_terms_init.json S4_terms_filtered.json

Paths are relative to the active workspace unless absolute or existing.

sf_terms.py 用于创建、转换、校验和比较术语文件。术语文件（S2_terms_init.json、
S4_terms_filtered.json、S5_terms_classified.json、S6_terms_defined.json、final/terms.json）
结构相同，逐阶段向前传递：{"stage", "scheme_id", "terms": [...]}。每个术语有稳定的
id T###、中英文名称、状态（candidate/kept/dropped/defined）、S5 赋予的 role 与 phase、
提取来源的原文引用、定义、辅助术语所 `supports` 的关键术语 id、出现的总结章节、来源、
丢弃原因、备注，以及仅存于 json 的 history[] 变更记录。术语永不删除，只标记 dropped，
以便 S4/S6 的决定可追溯。

.xlsx 副本是给人编辑的界面（一行一个术语，列表单元格用 "; " 连接，另有 legend 表解释列
和分类取值）。`to-xlsx` 生成它；`from-xlsx` 把修改合并回 json，追加 history 并为没有 id
的新行分配 id。json 为准。
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sf_common import (LIST_SEP, SUMMARY_SECTIONS, TERM_LIMIT_HARD, TERM_LIMIT_SOFT, TERM_LIST_COLUMNS,
                       TERM_ORIGINS, TERM_PHASES, TERM_ROLES, TERM_STATUSES, TERM_XLSX_COLUMNS, fail,
                       join_list, load_json, next_id, now_iso, relpath, resolve_workspace, save_json,
                       say, split_list)

COLUMN_DOC: dict[str, tuple[str, str]] = {
    "id": ("Stable id T###; leave empty for a new row and from-xlsx will assign one", "稳定 id；新行留空，from-xlsx 自动分配"),
    "name_zh": ("Canonical Chinese name; must match the **bold** text in the summary for callbacks", "中文规范名；回调术语需与总结中的加粗文字一致"),
    "name_en": ("snake_case English identifier; becomes the callback/function name", "snake_case 英文标识符，用作回调函数名"),
    "aliases": ("Other spellings found in the text, '; ' separated", "原文中的其他写法，用 '; ' 分隔"),
    "status": ("candidate | kept | dropped | defined", "候选 | 保留 | 丢弃 | 已定义"),
    "role": ("callback | parameter | constraint | aux_note | scope (S5)", "决策回调 | 参数 | 约束 | 辅助说明 | 适用范围（S5 赋值）"),
    "phase": ("market_context | entry | position_mgmt | exit | risk | philosophy (S5)", "环境/方向 | 入场 | 持仓管理 | 出场 | 风控 | 理念（S5 赋值）"),
    "is_key": ("Derived: TRUE when role == callback; downstream must implement it", "派生列：role 为 callback 时为 TRUE，下游必须实现"),
    "source_quote": ("Verbatim quote from S1_strategy_raw.md that motivated the term", "S1_strategy_raw.md 中的原文引用"),
    "definition_zh": ("Agreed Chinese definition (S6)", "S6 中与用户商定的中文定义"),
    "definition_en": ("Agreed English definition (S6); the downstream agent reads this", "S6 中商定的英文定义，下游 agent 读取此列"),
    "supports": ("For aux/constraint/parameter terms: ids of the key terms this one explains", "辅助/约束/参数术语所说明的关键术语 id"),
    "appears_in": ("Summary sections: " + ", ".join(SUMMARY_SECTIONS), "出现在总结的哪些章节"),
    "origin": (" | ".join(TERM_ORIGINS), "术语来源"),
    "drop_reason": ("Why it was dropped in S4/S6, citing dialog ids like D003", "丢弃原因，引用对话 id"),
    "notes": ("Free notes", "自由备注"),
}


# --------------------------------------------------------------------------- #
# Basic structure
# --------------------------------------------------------------------------- #


def ws_path(ws_arg: str | None, p: str) -> Path:
    """
    Resolve a file argument: absolute or existing paths are used as-is,
    otherwise the name is looked up inside the active workspace.

    解析文件参数：绝对路径或已存在的路径原样使用，否则在活动 workspace 内查找。
    """
    path = Path(p)
    if path.is_absolute() or path.exists():
        return path
    return resolve_workspace(ws_arg) / p


def empty_file(stage: str) -> dict:
    """
    Return an empty term file structure for a stage.

    返回某阶段的空术语文件结构。
    """
    return {
        "_comment_en": "Term set. Source of truth; the .xlsx twin is generated by sf_terms.py to-xlsx and merged back by from-xlsx. Never delete a term: set status=dropped.",
        "_comment_zh": "术语集。此文件为真值；同名 .xlsx 由 sf_terms.py to-xlsx 生成，用 from-xlsx 合并回来。不要删除术语，改 status=dropped。",
        "stage": stage,
        "scheme_id": None,
        "terms": [],
    }


def blank_term(tid: str) -> dict:
    """
    Return a term dict with every field present and empty, so hand-written
    JSON and xlsx rows share one shape.

    返回字段齐全、值为空的术语字典，使手写 JSON 与 xlsx 行保持同一形状。
    """
    return {
        "id": tid, "name_zh": "", "name_en": "", "aliases": [], "status": "candidate",
        "role": None, "phase": None, "is_key": False, "source_quote": "", "definition_zh": "",
        "definition_en": "", "supports": [], "appears_in": [], "origin": "S2_extract",
        "drop_reason": "", "notes": "", "history": [],
    }


def normalise(term: dict) -> dict:
    """
    Fill missing fields, coerce list fields, derive is_key from role, and
    return the term. Called on every load so scripts tolerate partially
    written entries from subagents.

    补齐缺失字段、规整列表字段、由 role 派生 is_key 后返回术语。每次加载都会调用，使脚本能
    容忍子代理写出的不完整条目。
    """
    base = blank_term(term.get("id", ""))
    base.update({k: v for k, v in term.items() if v is not None or k in ("role", "phase")})
    for col in TERM_LIST_COLUMNS:
        base[col] = split_list(base.get(col))
    base["is_key"] = base.get("role") == "callback"
    for k in ("name_zh", "name_en", "source_quote", "definition_zh", "definition_en", "drop_reason", "notes"):
        base[k] = "" if base.get(k) is None else str(base[k]).strip()
    if not isinstance(base.get("history"), list):
        base["history"] = []
    return base


def load_terms(path: Path) -> dict:
    """
    Load and normalise a term file.

    加载并规整术语文件。
    """
    data = load_json(path)
    if "terms" not in data:
        fail(f"{relpath(path)} has no 'terms' list", "文件缺少 terms 列表")
    data["terms"] = [normalise(t) for t in data["terms"]]
    return data


def history(term: dict, by: str, field: str, old, new) -> None:
    """
    Append a change record to term.history (json only).

    向 term.history 追加一条变更记录（仅 json 保存）。
    """
    term.setdefault("history", []).append({"at": now_iso(), "by": by, "field": field, "from": old, "to": new})


# --------------------------------------------------------------------------- #
# Validation
# --------------------------------------------------------------------------- #


def load_category_values(cat_path: Path | None) -> tuple[set[str], set[str]]:
    """
    Return (role_values, phase_values) allowed by a categories file, or the
    defaults from sf_common when none is given. A scheme without a phase
    axis returns an empty phase set, which disables phase checks.

    返回分类文件允许的 (role 取值, phase 取值)；未提供文件时用 sf_common 的默认值。没有
    phase 轴的方案返回空集合，从而跳过 phase 校验。
    """
    if cat_path is None:
        return set(TERM_ROLES), set(TERM_PHASES)
    cats = load_json(cat_path)
    roles: set[str] = set()
    phases: set[str] = set()
    for axis in cats.get("axes", []):
        vals = {v["id"] for v in axis.get("values", [])}
        if axis["axis_id"] == "role":
            roles = vals
        elif axis["axis_id"] == "phase":
            phases = vals
    return roles, phases


def validate(data: dict, cat_path: Path | None = None, require_classified: bool = False,
             require_defined: bool = False) -> tuple[list[str], list[str]]:
    """
    Return (errors, warnings) for a term file. Errors: duplicate/malformed
    ids, duplicate names among live terms, bad enum values, dangling or
    dropped `supports` targets, more than TERM_LIMIT_HARD live terms, and,
    when required, live terms lacking role/phase or definitions. Warnings:
    more than TERM_LIMIT_SOFT live terms, aux/constraint terms with no
    `supports`, callbacks without an English definition, unknown sections.

    返回术语文件的 (错误, 警告)。错误：id 重复或格式错误、存活术语名称重复、枚举值非法、
    supports 指向不存在或已丢弃的术语、存活术语超过硬上限、以及按需检查的缺少 role/phase
    或定义。警告：存活术语超过软上限、辅助/约束术语没有 supports、回调缺英文定义、章节名
    未知。
    """
    errors: list[str] = []
    warnings: list[str] = []
    roles, phases = load_category_values(cat_path)
    terms = data["terms"]
    ids = [t["id"] for t in terms]
    for tid in ids:
        if not re.fullmatch(r"T\d{3,}", tid):
            errors.append(f"malformed id {tid!r}")
    dup = {i for i in ids if ids.count(i) > 1}
    if dup:
        errors.append(f"duplicate ids: {sorted(dup)}")
    live = [t for t in terms if t["status"] != "dropped"]
    live_ids = {t["id"] for t in live}
    for key in ("name_zh", "name_en"):
        names = [t[key] for t in live if t[key]]
        d = {n for n in names if names.count(n) > 1}
        if d:
            errors.append(f"duplicate {key} among live terms: {sorted(d)}")
    for t in terms:
        tag = f"{t['id']} {t['name_zh']}"
        if t["status"] not in TERM_STATUSES:
            errors.append(f"{tag}: bad status {t['status']!r}")
        if t["origin"] not in TERM_ORIGINS:
            errors.append(f"{tag}: bad origin {t['origin']!r}")
        if t["status"] == "dropped":
            if not t["drop_reason"]:
                warnings.append(f"{tag}: dropped without drop_reason")
            continue
        if not t["name_zh"]:
            errors.append(f"{t['id']}: empty name_zh")
        if t["name_en"] and not re.fullmatch(r"[a-z][a-z0-9_]*", t["name_en"]):
            errors.append(f"{tag}: name_en {t['name_en']!r} is not snake_case")
        if t["role"] is not None and t["role"] not in roles:
            errors.append(f"{tag}: role {t['role']!r} not in {sorted(roles)}")
        if phases and t["phase"] is not None and t["phase"] not in phases:
            errors.append(f"{tag}: phase {t['phase']!r} not in {sorted(phases)}")
        if require_classified:
            if t["role"] is None:
                errors.append(f"{tag}: role missing (required after S5)")
            if phases and t["phase"] is None:
                errors.append(f"{tag}: phase missing (required after S5)")
            if not t["name_en"]:
                errors.append(f"{tag}: name_en missing (required after S5)")
        if require_defined and not t["definition_zh"]:
            errors.append(f"{tag}: definition_zh missing (required after S6)")
        if require_defined and t["role"] == "callback" and not t["definition_en"]:
            warnings.append(f"{tag}: callback without definition_en (downstream reads English)")
        for s in t["supports"]:
            if s not in {x["id"] for x in terms}:
                errors.append(f"{tag}: supports unknown id {s}")
            elif s not in live_ids:
                errors.append(f"{tag}: supports dropped term {s}")
        if t["role"] in ("aux_note", "constraint", "parameter") and not t["supports"]:
            warnings.append(f"{tag}: {t['role']} term with empty supports (which key term does it explain?)")
        for sec in t["appears_in"]:
            if sec not in SUMMARY_SECTIONS:
                warnings.append(f"{tag}: unknown section {sec!r}")
    n = len(live)
    if n > TERM_LIMIT_HARD:
        errors.append(f"{n} live terms exceeds hard limit {TERM_LIMIT_HARD}")
    elif n > TERM_LIMIT_SOFT:
        warnings.append(f"{n} live terms exceeds soft limit {TERM_LIMIT_SOFT}; consider merging")
    return errors, warnings


# --------------------------------------------------------------------------- #
# Excel round-trip
# --------------------------------------------------------------------------- #


def to_xlsx(data: dict, out: Path, cat_path: Path | None) -> None:
    """
    Write the term file as a workbook: sheet "terms" (one row per term,
    header = TERM_XLSX_COLUMNS, callbacks highlighted) and sheet "legend"
    (bilingual column explanations plus category values when a categories
    file is given). history[] is intentionally not exported.

    把术语文件写成工作簿："terms" 表一行一个术语（表头为 TERM_XLSX_COLUMNS，回调行高亮），
    "legend" 表给出双语列说明及分类取值（若提供分类文件）。history[] 有意不导出。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "terms"
    ws.append(TERM_XLSX_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    key_fill = PatternFill("solid", fgColor="FFF2CC")
    drop_fill = PatternFill("solid", fgColor="E7E6E6")
    for t in data["terms"]:
        row = []
        for col in TERM_XLSX_COLUMNS:
            v = t.get(col)
            if col in TERM_LIST_COLUMNS:
                v = join_list(v)
            elif col == "is_key":
                v = bool(v)
            row.append(v)
        ws.append(row)
        r = ws.max_row
        if t["status"] == "dropped":
            for c in ws[r]:
                c.fill = drop_fill
        elif t["is_key"]:
            for c in ws[r]:
                c.fill = key_fill
    widths = {"id": 7, "name_zh": 18, "name_en": 26, "aliases": 18, "status": 10, "role": 12, "phase": 15,
              "is_key": 7, "source_quote": 50, "definition_zh": 45, "definition_en": 45, "supports": 14,
              "appears_in": 22, "origin": 12, "drop_reason": 30, "notes": 30}
    for i, col in enumerate(TERM_XLSX_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 15)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"

    lg = wb.create_sheet("legend")
    lg.append(["column", "meaning (EN)", "含义 (中文)"])
    for col in TERM_XLSX_COLUMNS:
        en, zh = COLUMN_DOC[col]
        lg.append([col, en, zh])
    lg.append([])
    lg.append(["Agents read the json, not this workbook. Edit cells here, then run sf_terms.py from-xlsx.",
               "", "本表供人编辑；改完后运行 sf_terms.py from-xlsx 合并回 json。"])
    if cat_path is not None:
        cats = load_json(cat_path)
        lg.append([])
        lg.append([f"categories scheme: {cats.get('scheme_id')}", cats.get("name_en", ""), cats.get("name_zh", "")])
        for axis in cats.get("axes", []):
            lg.append([f"axis {axis['axis_id']}", axis.get("name_en", ""), axis.get("name_zh", "")])
            for v in axis.get("values", []):
                flag = "  [needs wrapper]" if v.get("needs_wrapper") else ""
                lg.append([f"  {v['id']}", v.get("name_en", "") + flag, v.get("name_zh", "")])
    for col, w in (("A", 24), ("B", 80), ("C", 60)):
        lg.column_dimensions[col].width = w
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def from_xlsx(data: dict, xlsx: Path, by: str = "user_edit") -> int:
    """
    Merge the "terms" sheet into the term data: rows with an id update that
    term's flat columns (history entries record each change), rows without an
    id become new terms with fresh ids and origin=user_edit, rows missing
    from the sheet are kept and reported. Returns the number of changed or
    added terms.

    把 "terms" 表合并进术语数据：有 id 的行更新对应术语的扁平列（每处改动记入 history），
    无 id 的行成为新术语（分配新 id，origin=user_edit），表中缺失的行保留并报告。返回改动
    或新增的术语数。
    """
    wb = load_workbook(xlsx, data_only=True)
    if "terms" not in wb.sheetnames:
        fail(f"{relpath(xlsx)} has no 'terms' sheet", "工作簿缺少 terms 表")
    sh = wb["terms"]
    rows = list(sh.iter_rows(values_only=True))
    if not rows:
        return 0
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    unknown = [h for h in header if h and h not in TERM_XLSX_COLUMNS]
    if unknown:
        say(f"ignoring unknown columns {unknown}", "忽略未知列")
    by_id = {t["id"]: t for t in data["terms"]}
    seen: set[str] = set()
    changed = 0
    for raw in rows[1:]:
        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
            continue
        cells = dict(zip(header, raw))
        tid = str(cells.get("id") or "").strip()
        if tid and tid in by_id:
            term = by_id[tid]
        else:
            tid = next_id(by_id.keys(), "T")
            term = blank_term(tid)
            term["origin"] = "user_edit"
            data["terms"].append(term)
            by_id[tid] = term
            history(term, by, "created", None, "from xlsx row")
            changed += 1
        seen.add(tid)
        touched = False
        for col in TERM_XLSX_COLUMNS:
            if col in ("id", "is_key") or col not in cells:
                continue
            new = cells[col]
            if col in TERM_LIST_COLUMNS:
                new = split_list(new)
            elif col in ("role", "phase"):
                new = (str(new).strip() or None) if new is not None else None
            else:
                new = "" if new is None else str(new).strip()
            if col in ("origin", "status") and new == "":
                continue  # an empty enum cell means "unchanged", never "blank" / 空枚举单元格视为未修改
            old = term.get(col)
            if new != old:
                history(term, by, col, old, new)
                term[col] = new
                touched = True
        if touched:
            changed += 1
        term["is_key"] = term.get("role") == "callback"
    missing = [t["id"] for t in data["terms"] if t["id"] not in seen]
    if missing:
        say(f"rows absent from xlsx kept unchanged (set status=dropped to remove): {missing}",
            "xlsx 中缺失的行保持不变（要移除请把 status 设为 dropped）")
    return changed


# --------------------------------------------------------------------------- #
# Commands
# --------------------------------------------------------------------------- #


def cmd_new(args: argparse.Namespace) -> None:
    """
    Create an empty term file for a stage.

    为某阶段创建空术语文件。
    """
    out = ws_path(args.workspace, args.path)
    if out.exists() and not args.force:
        fail(f"{relpath(out)} exists (use --force)", "文件已存在")
    save_json(out, empty_file(args.stage))
    say(f"created {relpath(out)}", "已创建")


def cmd_add(args: argparse.Namespace) -> None:
    """
    Append one term from CLI flags (used when a dialog in S3/S6 introduces a
    new term) and print its id.

    通过命令行参数追加一个术语（S3/S6 对话中出现新术语时使用）并打印其 id。
    """
    path = ws_path(args.workspace, args.path)
    data = load_terms(path)
    tid = next_id((t["id"] for t in data["terms"]), "T")
    t = blank_term(tid)
    t.update({"name_zh": args.name_zh, "name_en": args.name_en or "", "status": args.status,
              "role": args.role, "phase": args.phase, "source_quote": args.source_quote or "",
              "definition_zh": args.definition_zh or "", "definition_en": args.definition_en or "",
              "supports": split_list(args.supports), "appears_in": split_list(args.appears_in),
              "origin": args.origin, "notes": args.notes or ""})
    t["is_key"] = t["role"] == "callback"
    history(t, args.origin, "created", None, args.name_zh)
    data["terms"].append(t)
    save_json(path, data)
    say(f"added {tid} {args.name_zh}", "已添加术语")


def cmd_carry(args: argparse.Namespace) -> None:
    """
    Copy a term file forward to the next stage's filename, stamping the new
    stage and a history entry on every term. The agent then edits the copy.

    把术语文件复制为下一阶段的文件名，写入新阶段名并在每个术语的 history 里记一笔。
    之后 agent 在副本上编辑。
    """
    src = ws_path(args.workspace, args.src)
    dst = ws_path(args.workspace, args.dst)
    if dst.exists() and not args.force:
        fail(f"{relpath(dst)} exists (use --force)", "目标文件已存在")
    data = load_terms(src)
    data["stage"] = args.stage
    for t in data["terms"]:
        history(t, args.stage, "carried", src.name, dst.name)
    save_json(dst, data)
    say(f"carried {relpath(src)} -> {relpath(dst)}", "已复制到下一阶段")


def cmd_to_xlsx(args: argparse.Namespace) -> None:
    """
    Render <file>.json to <file>.xlsx (or --out).

    把 json 渲染为同名 xlsx（或 --out 指定路径）。
    """
    src = ws_path(args.workspace, args.path)
    data = load_terms(src)
    out = Path(args.out) if args.out else src.with_suffix(".xlsx")
    cat = ws_path(args.workspace, args.categories) if args.categories else None
    to_xlsx(data, out, cat)
    say(f"wrote {relpath(out)} ({len(data['terms'])} terms)", "已写出 xlsx")


def cmd_from_xlsx(args: argparse.Namespace) -> None:
    """
    Merge <file>.xlsx back into <file>.json (or --out), then validate and
    re-render the xlsx so both twins match.

    把 xlsx 合并回同名 json（或 --out），然后校验并重新生成 xlsx，使两者一致。
    """
    xlsx = ws_path(args.workspace, args.path)
    target = Path(args.out) if args.out else xlsx.with_suffix(".json")
    data = load_terms(target) if target.exists() else empty_file(args.stage or "S2")
    n = from_xlsx(data, xlsx)
    save_json(target, data)
    errors, warnings = validate(data)
    for w in warnings:
        say(f"warning: {w}")
    for e in errors:
        say(f"ERROR: {e}")
    cat = ws_path(args.workspace, args.categories) if args.categories else None
    to_xlsx(data, xlsx, cat)
    say(f"merged {n} changed/added terms into {relpath(target)}", "已合并回 json")


def cmd_validate(args: argparse.Namespace) -> None:
    """
    Validate a term file and exit 1 on errors; --fix rewrites the file with
    normalised fields (derived is_key, trimmed strings).

    校验术语文件，有错误则以 1 退出；--fix 会用规整后的字段重写文件。
    """
    path = ws_path(args.workspace, args.path)
    data = load_terms(path)
    cat = ws_path(args.workspace, args.categories) if args.categories else None
    errors, warnings = validate(data, cat, args.require_classified, args.require_defined)
    live = [t for t in data["terms"] if t["status"] != "dropped"]
    say(f"{relpath(path)}: {len(data['terms'])} terms, {len(live)} live, "
        f"{sum(1 for t in live if t['is_key'])} callbacks", "术语统计")
    for w in warnings:
        say(f"warning: {w}")
    for e in errors:
        say(f"ERROR: {e}")
    if args.fix:
        save_json(path, data)
    if errors:
        raise SystemExit(1)
    say("valid", "校验通过")


def cmd_diff(args: argparse.Namespace) -> None:
    """
    Print terms added, dropped, or changed between two term files (compares
    flat columns only).

    打印两个术语文件之间新增、丢弃、修改的术语（只比较扁平列）。
    """
    a = load_terms(ws_path(args.workspace, args.a))
    b = load_terms(ws_path(args.workspace, args.b))
    am = {t["id"]: t for t in a["terms"]}
    bm = {t["id"]: t for t in b["terms"]}
    for tid in sorted(set(bm) - set(am)):
        print(f"+ {tid} {bm[tid]['name_zh']} ({bm[tid]['origin']})")
    for tid in sorted(set(am) & set(bm)):
        x, y = am[tid], bm[tid]
        if x["status"] != "dropped" and y["status"] == "dropped":
            print(f"- {tid} {x['name_zh']}: dropped ({y['drop_reason']})")
            continue
        diffs = [c for c in TERM_XLSX_COLUMNS if c != "is_key" and x.get(c) != y.get(c)]
        if diffs:
            print(f"~ {tid} {y['name_zh']}: {', '.join(diffs)}")
    for tid in sorted(set(am) - set(bm)):
        print(f"! {tid} {am[tid]['name_zh']}: missing in second file (terms must not be deleted)")


def build_parser() -> argparse.ArgumentParser:
    """
    Build the CLI.

    构建命令行。
    """
    p = argparse.ArgumentParser(description="term files / 术语文件")
    p.add_argument("--workspace", "-w")
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("new")
    s.add_argument("path")
    s.add_argument("--stage", required=True)
    s.add_argument("--force", action="store_true")
    s.set_defaults(func=cmd_new)

    s = sub.add_parser("add")
    s.add_argument("path")
    s.add_argument("--name-zh", required=True)
    s.add_argument("--name-en")
    s.add_argument("--status", default="kept", choices=TERM_STATUSES)
    s.add_argument("--role", choices=TERM_ROLES)
    s.add_argument("--phase", choices=TERM_PHASES)
    s.add_argument("--source-quote")
    s.add_argument("--definition-zh")
    s.add_argument("--definition-en")
    s.add_argument("--supports", help="'; ' separated term ids")
    s.add_argument("--appears-in", help="'; ' separated section keys")
    s.add_argument("--origin", default="S6_dialog", choices=TERM_ORIGINS)
    s.add_argument("--notes")
    s.set_defaults(func=cmd_add)

    s = sub.add_parser("carry")
    s.add_argument("src")
    s.add_argument("dst")
    s.add_argument("--stage", required=True)
    s.add_argument("--force", action="store_true")
    s.set_defaults(func=cmd_carry)

    s = sub.add_parser("to-xlsx")
    s.add_argument("path")
    s.add_argument("--out")
    s.add_argument("--categories")
    s.set_defaults(func=cmd_to_xlsx)

    s = sub.add_parser("from-xlsx")
    s.add_argument("path")
    s.add_argument("--out")
    s.add_argument("--stage")
    s.add_argument("--categories")
    s.set_defaults(func=cmd_from_xlsx)

    s = sub.add_parser("validate")
    s.add_argument("path")
    s.add_argument("--categories")
    s.add_argument("--require-classified", action="store_true")
    s.add_argument("--require-defined", action="store_true")
    s.add_argument("--fix", action="store_true")
    s.set_defaults(func=cmd_validate)

    s = sub.add_parser("diff")
    s.add_argument("a")
    s.add_argument("b")
    s.set_defaults(func=cmd_diff)
    return p


def main(argv: list[str] | None = None) -> None:
    """
    CLI entry point.

    命令行入口。
    """
    args = build_parser().parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main()
