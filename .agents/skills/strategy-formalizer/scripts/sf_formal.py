#!/usr/bin/env python3
"""
Agents: read the English part of every docstring only. 中文段落仅供人类阅读。

sf_formal.py — S7 formalization: callbacks, parameters, relations and the
Python stub that the downstream trading agent fills in.

S7_formal.json is derived from the defined term file (S6_terms_defined.json)
and holds three lists. `callbacks` has one entry per role=callback term: the
function name (the term's name_en), where in the summary it is invoked, its
inputs, its output (usually an enum such as long/short/uncertain), and the
ids of the auxiliary, constraint and parameter terms the implementer must
read. `parameters` has one entry per role=parameter term with default/unit/
range. `relations` are typed edges between terms (supports, constrains,
parameterizes, triggers, sequence, conflicts) and feed the S7 graph.
`scaffold` builds a first version mechanically from the term file (the agent
then fills inputs/outputs/descriptions), `validate` checks it against the
terms, `to-xlsx`/`from-xlsx` give the human edit surface, and `gen-stub`
renders S7_callbacks_stub.py: an abstract class with one method per callback
whose docstring carries the definition and every auxiliary note, so the
downstream agent has everything in one place.

    sf_formal.py scaffold  S6_terms_defined.json S7_formal.json
    sf_formal.py validate  S7_formal.json --terms S6_terms_defined.json
    sf_formal.py to-xlsx   S7_formal.json
    sf_formal.py from-xlsx S7_formal.xlsx
    sf_formal.py gen-stub  S7_formal.json --terms S6_terms_defined.json --out S7_callbacks_stub.py

sf_formal.py 负责 S7 形式化：回调、参数、关系，以及供下游交易 agent 填充的 Python 桩
代码。S7_formal.json 由已定义的术语文件（S6_terms_defined.json）派生，包含三个列表：
`callbacks` 对应每个 role=callback 的术语：函数名（术语的 name_en）、在总结中的调用位置、
输入、输出（通常是 多/空/不确定 这类枚举），以及实现者必须阅读的辅助、约束、参数术语 id；
`parameters` 对应每个 role=parameter 的术语，含默认值/单位/范围；`relations` 是术语之间
带类型的边（supports、constrains、parameterizes、triggers、sequence、conflicts），也是 S7
关系图的数据来源。`scaffold` 从术语文件机械地生成初版（之后由 agent 补全输入/输出/描述），
`validate` 对照术语文件校验，`to-xlsx`/`from-xlsx` 提供人工编辑界面，`gen-stub` 生成
S7_callbacks_stub.py：一个抽象类，每个回调一个方法，docstring 里带有定义和全部辅助说明，
让下游 agent 在一处看到所需的一切。
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import networkx as nx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from sf_common import (RELATION_TYPES, TEMPLATES_DIR, fail, join_list, load_json, next_id, now_iso, read_text,
                       relpath, resolve_workspace, save_json, say, split_list)
from sf_terms import load_terms

ROLE_TO_RELATION = {"aux_note": "supports", "constraint": "constrains", "parameter": "parameterizes",
                    "scope": "supports"}

CB_COLUMNS = ["id", "term_id", "name_en", "name_zh", "invoked_in", "inputs", "output_type", "output_values",
              "aux_term_ids", "constraint_term_ids", "parameter_term_ids", "description_en", "description_zh"]
PARAM_COLUMNS = ["id", "term_id", "name_en", "name_zh", "default", "unit", "range", "description_en", "description_zh"]
REL_COLUMNS = ["from", "to", "type", "note"]


def ws_path(ws_arg: str | None, p: str) -> Path:
    """
    Resolve a file argument against the active workspace (see sf_terms).

    相对活动 workspace 解析文件参数（同 sf_terms）。
    """
    path = Path(p)
    if path.is_absolute() or path.exists():
        return path
    return resolve_workspace(ws_arg) / p


def empty_formal(terms_file: str) -> dict:
    """
    Return an empty S7_formal.json structure.

    返回空的 S7_formal.json 结构。
    """
    return {
        "_comment_en": "Formalized strategy interface for the downstream trading agent. callbacks = decisions the downstream agent must implement; parameters = tunables; relations = typed edges between terms (drawn by sf_graph.py). Source of truth; the .xlsx twin is generated.",
        "_comment_zh": "面向下游交易 agent 的形式化策略接口。callbacks = 下游必须实现的决策；parameters = 可调参数；relations = 术语之间的带类型的边（由 sf_graph.py 绘图）。此文件为真值，xlsx 由脚本生成。",
        "stage": "S7",
        "terms_file": terms_file,
        "generated_at": now_iso(),
        "callbacks": [],
        "parameters": [],
        "relations": [],
    }


# --------------------------------------------------------------------------- #
# scaffold
# --------------------------------------------------------------------------- #


def scaffold(terms: dict, terms_file: str) -> dict:
    """
    Build a first S7_formal.json from a defined term file: one callback per
    live role=callback term (inputs empty, output enum empty, descriptions
    copied from the definitions), one parameter per role=parameter term, and
    one relation per `supports` link typed by the supporting term's role.
    Each callback also collects the ids of aux/constraint/parameter terms
    that support it. The agent completes inputs/outputs afterwards.

    从已定义的术语文件生成初版 S7_formal.json：每个存活的 role=callback 术语生成一个回调
    （输入为空、输出枚举为空、描述从定义复制），每个 role=parameter 术语生成一个参数，每条
    `supports` 链接生成一条关系（类型由提供支持的术语的 role 决定）。每个回调同时收集支持它
    的辅助/约束/参数术语 id。之后由 agent 补全输入/输出。
    """
    live = [t for t in terms["terms"] if t["status"] != "dropped"]
    formal = empty_formal(terms_file)
    cb_by_term: dict[str, dict] = {}
    for t in live:
        if t["role"] == "callback":
            cb = {
                "id": next_id((c["id"] for c in formal["callbacks"]), "CB", 2),
                "term_id": t["id"],
                "name_en": t["name_en"],
                "name_zh": t["name_zh"],
                "invoked_in": (t["appears_in"][0] if t["appears_in"] else "") ,
                "inputs": [],
                "output": {"type": "enum", "values": [], "description_en": ""},
                "aux_term_ids": [],
                "constraint_term_ids": [],
                "parameter_term_ids": [],
                "description_en": t["definition_en"],
                "description_zh": t["definition_zh"],
            }
            formal["callbacks"].append(cb)
            cb_by_term[t["id"]] = cb
    for t in live:
        if t["role"] == "parameter":
            formal["parameters"].append({
                "id": next_id((p["id"] for p in formal["parameters"]), "P", 2),
                "term_id": t["id"], "name_en": t["name_en"], "name_zh": t["name_zh"],
                "default": "", "unit": "", "range": "",
                "description_en": t["definition_en"], "description_zh": t["definition_zh"],
            })
    for t in live:
        rel_type = ROLE_TO_RELATION.get(t["role"] or "")
        for target in t["supports"]:
            if rel_type:
                formal["relations"].append({"from": t["id"], "to": target, "type": rel_type, "note": ""})
            cb = cb_by_term.get(target)
            if cb is not None:
                bucket = {"aux_note": "aux_term_ids", "scope": "aux_term_ids", "constraint": "constraint_term_ids",
                          "parameter": "parameter_term_ids"}.get(t["role"] or "")
                if bucket and t["id"] not in cb[bucket]:
                    cb[bucket].append(t["id"])
    return formal


# --------------------------------------------------------------------------- #
# validate
# --------------------------------------------------------------------------- #


def validate(formal: dict, terms: dict) -> tuple[list[str], list[str]]:
    """
    Return (errors, warnings). Errors: callback pointing at a missing,
    dropped or non-callback term; a live callback term with no callback or
    with two; name_en mismatch; relation endpoints or types unknown; a cycle
    among `sequence` relations. Warnings: callbacks with empty inputs, empty
    output values or empty description_en; relations referencing dropped
    terms.

    返回 (错误, 警告)。错误：回调指向不存在、已丢弃或非 callback 的术语；存活的 callback
    术语没有回调或有两个；name_en 不一致；关系端点或类型未知；sequence 关系成环。警告：
    回调输入为空、输出取值为空或英文描述为空；关系引用了已丢弃术语。
    """
    errors: list[str] = []
    warnings: list[str] = []
    tmap = {t["id"]: t for t in terms["terms"]}
    live_cb = {t["id"] for t in terms["terms"] if t["status"] != "dropped" and t["role"] == "callback"}
    seen: dict[str, int] = {}
    for cb in formal["callbacks"]:
        tag = f"{cb['id']} {cb.get('name_en')}"
        t = tmap.get(cb["term_id"])
        if t is None:
            errors.append(f"{tag}: term {cb['term_id']} does not exist")
            continue
        if t["status"] == "dropped":
            errors.append(f"{tag}: term {cb['term_id']} is dropped")
        if t["role"] != "callback":
            errors.append(f"{tag}: term {cb['term_id']} has role {t['role']}, not callback")
        if cb.get("name_en") != t["name_en"]:
            errors.append(f"{tag}: name_en differs from term name_en {t['name_en']!r}")
        seen[cb["term_id"]] = seen.get(cb["term_id"], 0) + 1
        if not cb.get("inputs"):
            warnings.append(f"{tag}: no inputs declared")
        out = cb.get("output") or {}
        if out.get("type") == "enum" and not out.get("values"):
            warnings.append(f"{tag}: enum output has no values")
        if not cb.get("description_en"):
            warnings.append(f"{tag}: description_en empty (downstream reads English)")
        for bucket in ("aux_term_ids", "constraint_term_ids", "parameter_term_ids"):
            for tid in cb.get(bucket, []):
                if tid not in tmap:
                    errors.append(f"{tag}: {bucket} references unknown {tid}")
    for tid in live_cb:
        n = seen.get(tid, 0)
        if n == 0:
            errors.append(f"callback term {tid} {tmap[tid]['name_zh']} has no callback entry")
        elif n > 1:
            errors.append(f"callback term {tid} has {n} callback entries")
    for p in formal.get("parameters", []):
        if p["term_id"] not in tmap:
            errors.append(f"{p['id']}: term {p['term_id']} does not exist")
    known = set(tmap) | {c["id"] for c in formal["callbacks"]} | {p["id"] for p in formal.get("parameters", [])}
    g = nx.DiGraph()
    for r in formal.get("relations", []):
        if r["type"] not in RELATION_TYPES:
            errors.append(f"relation {r['from']}->{r['to']}: bad type {r['type']!r}")
        for end in ("from", "to"):
            if r[end] not in known:
                errors.append(f"relation {r['from']}->{r['to']}: unknown {end} {r[end]}")
            elif r[end] in tmap and tmap[r[end]]["status"] == "dropped":
                warnings.append(f"relation {r['from']}->{r['to']}: {end} is a dropped term")
        if r["type"] == "sequence":
            g.add_edge(r["from"], r["to"])
    try:
        cycle = nx.find_cycle(g)
        errors.append(f"sequence relations form a cycle: {cycle}")
    except nx.NetworkXNoCycle:
        pass
    return errors, warnings


# --------------------------------------------------------------------------- #
# xlsx
# --------------------------------------------------------------------------- #


def inputs_to_cell(inputs: list[dict]) -> str:
    """
    Render inputs as "name:type; name:type" for a spreadsheet cell.

    把输入列表渲染为 "name:type; name:type" 形式的单元格文本。
    """
    return join_list(f"{i.get('name')}:{i.get('type', '')}" for i in inputs)


def cell_to_inputs(cell) -> list[dict]:
    """
    Parse "name:type; name:type" back into input dicts (descriptions are
    kept only in json).

    把 "name:type; name:type" 解析回输入字典列表（描述仅保存在 json）。
    """
    out = []
    for item in split_list(cell):
        name, _, typ = item.partition(":")
        out.append({"name": name.strip(), "type": typ.strip(), "source_term_id": None, "description_en": ""})
    return out


def to_xlsx(formal: dict, out: Path) -> None:
    """
    Write sheets callbacks / parameters / relations / legend.

    写出 callbacks / parameters / relations / legend 四张表。
    """
    wb = Workbook()
    sh = wb.active
    sh.title = "callbacks"
    sh.append(CB_COLUMNS)
    for cb in formal["callbacks"]:
        out_ = cb.get("output") or {}
        sh.append([cb["id"], cb["term_id"], cb["name_en"], cb.get("name_zh", ""), cb.get("invoked_in", ""),
                   inputs_to_cell(cb.get("inputs", [])), out_.get("type", ""), join_list(out_.get("values", [])),
                   join_list(cb.get("aux_term_ids")), join_list(cb.get("constraint_term_ids")),
                   join_list(cb.get("parameter_term_ids")), cb.get("description_en", ""), cb.get("description_zh", "")])
    ps = wb.create_sheet("parameters")
    ps.append(PARAM_COLUMNS)
    for p in formal.get("parameters", []):
        ps.append([p.get(c, "") for c in PARAM_COLUMNS])
    rs = wb.create_sheet("relations")
    rs.append(REL_COLUMNS)
    for r in formal.get("relations", []):
        rs.append([r.get(c, "") for c in REL_COLUMNS])
    lg = wb.create_sheet("legend")
    lg.append(["sheet/column", "meaning (EN)", "含义 (中文)"])
    lg.append(["callbacks.inputs", "name:type; name:type", "输入，形如 name:type; name:type"])
    lg.append(["callbacks.output_type", "enum | number | bool | price | object", "输出类型"])
    lg.append(["callbacks.output_values", "enum values, '; ' separated, e.g. long; short; uncertain", "枚举取值，如 多; 空; 不确定 的英文"])
    lg.append(["callbacks.*_term_ids", "term ids the implementer must read", "实现者必须阅读的术语 id"])
    lg.append(["relations.type", " | ".join(RELATION_TYPES), "关系类型"])
    lg.append(["", "Agents read the json; edit here then run sf_formal.py from-xlsx.", "供人编辑，改完运行 from-xlsx"])
    for ws_ in (sh, ps, rs, lg):
        for c in ws_[1]:
            c.font = Font(bold=True)
        for row in ws_.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for col in ws_.columns:
            ws_.column_dimensions[col[0].column_letter].width = 24
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def from_xlsx(formal: dict, xlsx: Path) -> None:
    """
    Merge the three sheets back into the formal dict: callbacks and
    parameters are matched by id (new rows without id get ids), relations
    are replaced wholesale by the sheet contents.

    把三张表合并回 formal 字典：callbacks 和 parameters 按 id 匹配（无 id 的新行分配 id），
    relations 整体以表内容替换。
    """
    wb = load_workbook(xlsx, data_only=True)

    def rows(name: str) -> list[dict]:
        """
        Read one sheet as a list of {header: value} dicts, skipping empty rows; missing sheets yield an empty list.

        把一张表读成 {表头: 值} 字典列表，跳过空行；表不存在时返回空列表。
        """
        if name not in wb.sheetnames:
            return []
        it = list(wb[name].iter_rows(values_only=True))
        if not it:
            return []
        header = [str(h).strip() if h else "" for h in it[0]]
        return [dict(zip(header, r)) for r in it[1:] if r and any(v not in (None, "") for v in r)]

    cbs = {c["id"]: c for c in formal["callbacks"]}
    for r in rows("callbacks"):
        cid = str(r.get("id") or "").strip() or next_id(cbs.keys(), "CB", 2)
        cb = cbs.setdefault(cid, {"id": cid, "inputs": [], "output": {}, "aux_term_ids": [],
                                  "constraint_term_ids": [], "parameter_term_ids": []})
        cb["term_id"] = str(r.get("term_id") or "").strip()
        cb["name_en"] = str(r.get("name_en") or "").strip()
        cb["name_zh"] = str(r.get("name_zh") or "").strip()
        cb["invoked_in"] = str(r.get("invoked_in") or "").strip()
        if r.get("inputs") not in (None, ""):
            cb["inputs"] = cell_to_inputs(r.get("inputs"))
        out_ = cb.setdefault("output", {})
        out_["type"] = str(r.get("output_type") or out_.get("type") or "enum").strip()
        out_["values"] = split_list(r.get("output_values"))
        out_.setdefault("description_en", "")
        cb["aux_term_ids"] = split_list(r.get("aux_term_ids"))
        cb["constraint_term_ids"] = split_list(r.get("constraint_term_ids"))
        cb["parameter_term_ids"] = split_list(r.get("parameter_term_ids"))
        cb["description_en"] = str(r.get("description_en") or "").strip()
        cb["description_zh"] = str(r.get("description_zh") or "").strip()
    formal["callbacks"] = list(cbs.values())
    ps = {p["id"]: p for p in formal.get("parameters", [])}
    for r in rows("parameters"):
        pid = str(r.get("id") or "").strip() or next_id(ps.keys(), "P", 2)
        p = ps.setdefault(pid, {"id": pid})
        for c in PARAM_COLUMNS[1:]:
            v = r.get(c)
            p[c] = "" if v is None else (v if isinstance(v, (int, float)) else str(v).strip())
    formal["parameters"] = list(ps.values())
    rels = rows("relations")
    if rels:
        formal["relations"] = [{"from": str(r.get("from") or "").strip(), "to": str(r.get("to") or "").strip(),
                                "type": str(r.get("type") or "").strip(), "note": str(r.get("note") or "").strip()}
                               for r in rels]


# --------------------------------------------------------------------------- #
# stub generation
# --------------------------------------------------------------------------- #


STANDARD_TYPES = ("MarketData", "Bar", "Position", "Instrument", "Direction", "EntryDecision", "StopDistance")


def py_type(name: str) -> str:
    """
    Map a type name used in formal.json (inputs[].type / output.type) to a
    Python annotation in the stub: the standard interface types defined in
    templates/callbacks_stub.template.py are used verbatim, `enum` outputs
    become Literal[...] (handled by py_literal_type), scalar names map to
    builtins, `Literal[...]` and `Optional[...]` strings pass through, and
    anything unknown becomes Any.

    把 formal.json 中的类型名（inputs[].type / output.type）映射为桩代码的 Python 标注：
    templates/callbacks_stub.template.py 定义的标准接口类型原样使用，enum 输出由
    py_literal_type 变成 Literal[...]，标量名映射为内置类型，`Literal[...]`/`Optional[...]`
    字符串原样透传，未知类型为 Any。
    """
    name = (name or "").strip()
    if name in STANDARD_TYPES or name.startswith(("Literal[", "Optional[", "list[", "dict[")):
        return name
    return {"number": "float", "price": "float", "float": "float", "bool": "bool", "int": "int",
            "str": "str", "string": "str"}.get(name, "Any")


def py_literal_type(output: dict) -> str:
    """
    Map an output spec to a Python annotation: enum -> Literal[...], a
    standard interface type or scalar via py_type, anything else -> Any.

    把输出规格映射为 Python 类型标注：enum -> Literal[...]，标准接口类型或标量经 py_type
    映射，其他 -> Any。
    """
    typ = (output or {}).get("type", "")
    if typ == "enum" and output.get("values"):
        return "Literal[" + ", ".join(repr(v) for v in output["values"]) + "]"
    return py_type(typ)


def _wrap(text: str, indent: str, width: int = 88) -> list[str]:
    """
    Wrap a paragraph into docstring lines at the given indent.

    把段落按缩进折行成 docstring 行。
    """
    import textwrap
    return textwrap.wrap(text, width=width - len(indent), initial_indent=indent, subsequent_indent=indent) or [indent]


def gen_stub(formal: dict, terms: dict, strategy_name: str) -> str:
    """
    Render the abstract callbacks class. The module docstring comes from
    templates/callbacks_stub.template.py; each callback becomes an abstract
    method whose docstring lists (EN first, then ZH) the definition, the
    invocation point, inputs, output, and every auxiliary / constraint /
    parameter term with its definition, so the downstream implementer never
    has to open the term file.

    渲染抽象回调类。模块 docstring 来自 templates/callbacks_stub.template.py；每个回调
    变成一个抽象方法，其 docstring（先英后中）列出定义、调用位置、输入、输出，以及每个
    辅助/约束/参数术语及其定义，使下游实现者无需再打开术语文件。
    """
    tmap = {t["id"]: t for t in terms["terms"]}
    header = read_text(TEMPLATES_DIR / "callbacks_stub.template.py")
    header = header.replace("{strategy_name}", strategy_name).replace("{generated_at}", now_iso())
    lines = [header.rstrip("\n"), "", ""]
    lines.append("class StrategyCallbacks(ABC):")
    lines.append('    """')
    lines.append("    Decision callbacks the downstream trading agent must implement. Each method")
    lines.append("    corresponds to one role=callback term; read its docstring for the agreed")
    lines.append("    definition and the auxiliary notes that constrain the implementation.")
    lines.append("")
    lines.append("    下游交易 agent 必须实现的决策回调。每个方法对应一个 role=callback 的术语；")
    lines.append("    方法 docstring 里是商定的定义和约束实现的辅助说明。")
    lines.append('    """')
    lines.append("")
    if not formal["callbacks"]:
        lines.append("    pass")
    for cb in formal["callbacks"]:
        t = tmap.get(cb["term_id"], {})
        params = ["self"]
        for i in cb.get("inputs", []):
            params.append(f"{i['name']}: {py_type(i.get('type', ''))}")
        ret = py_literal_type(cb.get("output") or {})
        lines.append("    @abstractmethod")
        lines.append(f"    def {cb['name_en']}({', '.join(params)}) -> {ret}:")
        ind = "        "
        lines.append(f'{ind}"""')
        lines.append(f"{ind}[{cb['id']} / {cb['term_id']}] {cb.get('name_zh', '')}")
        lines.append("")
        lines += _wrap("Definition (EN): " + (cb.get("description_en") or t.get("definition_en") or "(missing)"), ind)
        lines.append(f"{ind}Invoked in: {cb.get('invoked_in') or '(unspecified)'}")
        for i in cb.get("inputs", []):
            lines += _wrap(f"Input {i['name']} ({i.get('type', '')}): {i.get('description_en', '')}", ind)
        out_ = cb.get("output") or {}
        lines += _wrap(f"Output ({out_.get('type', '')}): {', '.join(out_.get('values', []))} {out_.get('description_en', '')}".rstrip(), ind)
        for bucket, label in (("aux_term_ids", "Auxiliary"), ("constraint_term_ids", "Constraint"),
                              ("parameter_term_ids", "Parameter")):
            for tid in cb.get(bucket, []):
                at = tmap.get(tid, {})
                lines += _wrap(f"{label} {tid} {at.get('name_zh', '')} / {at.get('name_en', '')}: "
                               f"{at.get('definition_en') or at.get('definition_zh') or at.get('source_quote', '')}", ind)
        lines.append("")
        lines += _wrap("定义（中文）: " + (cb.get("description_zh") or t.get("definition_zh") or "（缺失）"), ind)
        for bucket, label in (("aux_term_ids", "辅助"), ("constraint_term_ids", "约束"), ("parameter_term_ids", "参数")):
            for tid in cb.get(bucket, []):
                at = tmap.get(tid, {})
                lines += _wrap(f"{label} {tid} {at.get('name_zh', '')}: {at.get('definition_zh') or at.get('source_quote', '')}", ind)
        lines.append(f'{ind}"""')
        lines.append(f"{ind}raise NotImplementedError")
        lines.append("")
    if formal.get("parameters"):
        lines.append("")
        lines.append("PARAMETERS: dict[str, dict[str, Any]] = {")
        for p in formal["parameters"]:
            lines.append(f"    {p['name_en']!r}: {{'term_id': {p['term_id']!r}, 'default': {p.get('default')!r}, "
                         f"'unit': {p.get('unit', '')!r}, 'range': {p.get('range', '')!r}, "
                         f"'description_en': {p.get('description_en', '')!r}}},")
        lines.append("}")
        lines.append('"""Tunable parameters (role=parameter terms). 可调参数（role=parameter 的术语）。"""')
    return "\n".join(lines) + "\n"


# --------------------------------------------------------------------------- #
# Commands
# --------------------------------------------------------------------------- #


def cmd_scaffold(args: argparse.Namespace) -> None:
    """
    Generate S7_formal.json (and its xlsx) from a term file.

    从术语文件生成 S7_formal.json（及其 xlsx）。
    """
    tpath = ws_path(args.workspace, args.terms)
    out = ws_path(args.workspace, args.out)
    if out.exists() and not args.force:
        fail(f"{relpath(out)} exists (use --force)", "文件已存在")
    terms = load_terms(tpath)
    formal = scaffold(terms, tpath.name)
    save_json(out, formal)
    to_xlsx(formal, out.with_suffix(".xlsx"))
    say(f"scaffolded {len(formal['callbacks'])} callbacks, {len(formal['parameters'])} parameters, "
        f"{len(formal['relations'])} relations -> {relpath(out)}", "已生成初版形式化文件")


def cmd_validate(args: argparse.Namespace) -> None:
    """
    Validate formal.json against the term file; exit 1 on errors.

    对照术语文件校验 formal.json；有错误则以 1 退出。
    """
    formal = load_json(ws_path(args.workspace, args.path))
    terms = load_terms(ws_path(args.workspace, args.terms))
    errors, warnings = validate(formal, terms)
    for w in warnings:
        say(f"warning: {w}")
    for e in errors:
        say(f"ERROR: {e}")
    if errors:
        raise SystemExit(1)
    say("valid", "校验通过")


def cmd_to_xlsx(args: argparse.Namespace) -> None:
    """
    Render formal.json to xlsx.

    把 formal.json 渲染为 xlsx。
    """
    src = ws_path(args.workspace, args.path)
    out = Path(args.out) if args.out else src.with_suffix(".xlsx")
    to_xlsx(load_json(src), out)
    say(f"wrote {relpath(out)}", "已写出 xlsx")


def cmd_from_xlsx(args: argparse.Namespace) -> None:
    """
    Merge the xlsx back into formal.json and re-render the xlsx.

    把 xlsx 合并回 formal.json 并重新生成 xlsx。
    """
    xlsx = ws_path(args.workspace, args.path)
    target = Path(args.out) if args.out else xlsx.with_suffix(".json")
    formal = load_json(target) if target.exists() else empty_formal("")
    from_xlsx(formal, xlsx)
    save_json(target, formal)
    to_xlsx(formal, xlsx)
    say(f"merged into {relpath(target)}", "已合并回 json")


def cmd_gen_stub(args: argparse.Namespace) -> None:
    """
    Write the abstract callbacks class to --out.

    把抽象回调类写到 --out。
    """
    formal = load_json(ws_path(args.workspace, args.path))
    terms = load_terms(ws_path(args.workspace, args.terms))
    out = ws_path(args.workspace, args.out)
    name = args.strategy_name or resolve_workspace(args.workspace).name
    code = gen_stub(formal, terms, name)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(code, encoding="utf-8")
    compile(code, str(out), "exec")
    say(f"wrote {relpath(out)} ({len(formal['callbacks'])} methods)", "已生成桩代码")


def build_parser() -> argparse.ArgumentParser:
    """
    Build the CLI.

    构建命令行。
    """
    p = argparse.ArgumentParser(description="S7 formalization / 形式化")
    p.add_argument("--workspace", "-w")
    sub = p.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("scaffold")
    s.add_argument("terms")
    s.add_argument("out")
    s.add_argument("--force", action="store_true")
    s.set_defaults(func=cmd_scaffold)
    s = sub.add_parser("validate")
    s.add_argument("path")
    s.add_argument("--terms", required=True)
    s.set_defaults(func=cmd_validate)
    s = sub.add_parser("to-xlsx")
    s.add_argument("path")
    s.add_argument("--out")
    s.set_defaults(func=cmd_to_xlsx)
    s = sub.add_parser("from-xlsx")
    s.add_argument("path")
    s.add_argument("--out")
    s.set_defaults(func=cmd_from_xlsx)
    s = sub.add_parser("gen-stub")
    s.add_argument("path")
    s.add_argument("--terms", required=True)
    s.add_argument("--out", default="S7_callbacks_stub.py")
    s.add_argument("--strategy-name")
    s.set_defaults(func=cmd_gen_stub)
    return p


def main(argv: list[str] | None = None) -> None:
    """
    CLI entry point.

    命令行入口。
    """
    args = build_parser().parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main()
