#!/usr/bin/env python3
"""
Agents: read the English part of every docstring only. 中文段落仅供人类阅读。

sf_common.py — shared helpers for every strategy-formalizer script.

This module owns the things that all stage scripts agree on: where the repo
root and the active workspace are, the ordered list of stages S1..S8 with the
files each stage must produce, the enumerations used inside term.json /
dialog.json / formal.json, and small utilities (JSON I/O that keeps Chinese
readable, SHA-256 hashing for checkpoints, ISO timestamps, bilingual console
output). No script should hard-code a stage name or an output filename; it
should import them from here so the SKILL.md manifest and the code cannot
drift apart.

sf_common.py 是所有 strategy-formalizer 脚本共用的基础模块。它统一定义：仓库根目录与
当前活动 workspace 的定位方式、S1..S8 各阶段的顺序及每个阶段必须产出的文件清单、
term.json / dialog.json / formal.json 里用到的枚举值，以及一些小工具（保留中文的 JSON
读写、用于检查点的 SHA-256 哈希、ISO 时间戳、中英双语的控制台输出）。任何脚本都不应
自行硬编码阶段名或输出文件名，而应从这里导入，以保证 SKILL.md 的文件清单与代码一致。
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

# --------------------------------------------------------------------------- #
# Paths
# --------------------------------------------------------------------------- #

SCRIPTS_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPTS_DIR.parent                     # .agents/skills/strategy-formalizer
REPO_ROOT = SKILL_DIR.parents[2]                   # .agents/skills/<skill> -> repo root
WORKSPACE_ROOT = REPO_ROOT / "workspace"
ACTIVE_FILE = WORKSPACE_ROOT / "ACTIVE"            # holds the name of the active workspace
TEMPLATES_DIR = SKILL_DIR / "templates"
SCHEMAS_DIR = SKILL_DIR / "schemas"
STAGES_DIR = SKILL_DIR / "stages"

# --------------------------------------------------------------------------- #
# Stage definitions (single source of truth for SKILL.md's manifest)
# --------------------------------------------------------------------------- #

STAGES: list[str] = ["S1", "S2", "S3", "S4", "S5", "S6", "S7", "S8"]

STAGE_TITLES: dict[str, tuple[str, str]] = {
    "S1": ("Correct raw text", "矫正原始文本"),
    "S2": ("Summarize + extract terms (two subagents)", "总结 + 关键术语提取（两个子代理）"),
    "S3": ("Summary dialog with the user", "与用户多轮对话修订策略总结"),
    "S4": ("Filter / correct terms using S3 context", "依据 S3 对话过滤、矫正术语"),
    "S5": ("Classify terms, mark callbacks in bold", "术语分类并在总结中加粗回调术语"),
    "S6": ("Define every term with the user", "与用户多轮对话明确每个术语"),
    "S7": ("Formalize: callbacks, relations, graph, stub", "形式化：回调、关系、图、桩代码"),
    "S8": ("Double check terms vs summary, freeze final/", "术语与总结交叉核对并冻结 final/"),
}

STAGE_DOCS: dict[str, str] = {
    "S1": "S1_correct.md",
    "S2": "S2_summarize_and_extract.md",
    "S3": "S3_summary_dialog.md",
    "S4": "S4_term_filter.md",
    "S5": "S5_term_classify.md",
    "S6": "S6_term_define_dialog.md",
    "S7": "S7_formalize.md",
    "S8": "S8_double_check.md",
}

# (filename, required) pairs. Optional files are hashed if present.
STAGE_OUTPUTS: dict[str, list[tuple[str, bool]]] = {
    "S1": [("S1_sources.json", True), ("S1_strategy_raw.md", True), ("S1_corrections.md", True),
           ("S1_strategy_clean.md", True)],
    "S2": [("S2_summary_init.md", True), ("S2_terms_init.json", True), ("S2_terms_init.xlsx", True)],
    "S3": [("S3_dialog.json", True), ("S3_dialog.md", True), ("S3_summary_corrected.md", True)],
    "S4": [("S4_terms_filtered.json", True), ("S4_terms_filtered.xlsx", True)],
    "S5": [("S5_categories.json", True), ("S5_terms_classified.json", True),
           ("S5_terms_classified.xlsx", True), ("S5_summary_marked.md", True)],
    "S6": [("S6_dialog.json", True), ("S6_dialog.md", True), ("S6_terms_defined.json", True),
           ("S6_terms_defined.xlsx", True), ("S6_conflict_log.md", True)],
    "S7": [("S7_formal.json", True), ("S7_formal.xlsx", True), ("S7_term_graph.mmd", True),
           ("S7_term_graph.dot", True), ("S7_term_graph.pdf", False), ("S7_callbacks_stub.py", True),
           ("S7_strategy_driver.py", True)],
    "S8": [("S8_check_report.md", True), ("final/strategy.md", True), ("final/terms.json", True),
           ("final/formal.json", True), ("final/callbacks_stub.py", True), ("final/strategy_driver.py", True)],
}

# Which stage's term file is the "latest terms" input for a later stage.
TERM_FILE_BY_STAGE: dict[str, str] = {
    "S2": "S2_terms_init.json",
    "S4": "S4_terms_filtered.json",
    "S5": "S5_terms_classified.json",
    "S6": "S6_terms_defined.json",
    "S8": "final/terms.json",
}

DIALOG_STAGES = ("S3", "S6")

STAGE_STATUSES = ("pending", "in_progress", "awaiting_review", "done")

# --------------------------------------------------------------------------- #
# Enumerations shared by schemas, validators and converters
# --------------------------------------------------------------------------- #

TERM_STATUSES = ("candidate", "kept", "dropped", "defined")
TERM_ROLES = ("callback", "parameter", "constraint", "aux_note", "scope")
TERM_PHASES = ("market_context", "entry", "position_mgmt", "exit", "risk", "philosophy")
TERM_ORIGINS = ("S2_extract", "S3_dialog", "S4_filter", "S5_classify", "S6_dialog", "user_edit")
DIALOG_STATUSES = ("open", "answered", "closed")
RELATION_TYPES = ("supports", "constrains", "parameterizes", "triggers", "sequence", "conflicts")
SUMMARY_SECTIONS = ("summary", "one_liner", "premise", "execution_steps", "exit_and_risk", "scope_and_params")
SUMMARY_SECTION_TITLES: dict[str, tuple[str, str]] = {
    "summary": ("Strategy summary", "策略总结"),
    "one_liner": ("One-line overview", "一句话概述"),
    "premise": ("Premise and philosophy", "策略前提与理念"),
    "execution_steps": ("Execution steps", "执行步骤"),
    "exit_and_risk": ("Exit and risk control", "出场与风控"),
    "scope_and_params": ("Scope and parameters", "适用范围与参数"),
}

TERM_LIMIT_SOFT = 30
TERM_LIMIT_HARD = 50

# Flat columns that round-trip between term.json and term.xlsx (history[] does not).
TERM_XLSX_COLUMNS: list[str] = [
    "id", "name_zh", "name_en", "aliases", "status", "role", "phase", "is_key",
    "source_quote", "definition_zh", "definition_en", "supports", "appears_in",
    "origin", "drop_reason", "notes",
]
TERM_LIST_COLUMNS = ("aliases", "supports", "appears_in")
LIST_SEP = "; "

# --------------------------------------------------------------------------- #
# Small utilities
# --------------------------------------------------------------------------- #


def now_iso() -> str:
    """
    Return the current local time as an ISO-8601 string with timezone, second
    precision. Used for every timestamp written into workspace files so that
    dialog logs and state.json are comparable.

    返回带时区、精确到秒的 ISO-8601 当前时间字符串。workspace 内所有时间戳都用它生成，
    便于对话记录和 state.json 之间相互比较。
    """
    return datetime.now(timezone.utc).astimezone().replace(microsecond=0).isoformat()


def load_json(path: Path | str) -> Any:
    """
    Read a UTF-8 JSON file and return the parsed object. Raises a clear error
    naming the file when it is missing or malformed, because scripts are often
    run by the agent right after the user hand-edited the file.

    读取 UTF-8 编码的 JSON 文件并返回解析结果。文件缺失或格式错误时抛出带文件名的清晰
    错误，因为脚本经常在用户手工编辑文件之后由 agent 立即运行。
    """
    p = Path(path)
    try:
        with p.open("r", encoding="utf-8") as fh:
            return json.load(fh)
    except FileNotFoundError:
        raise SystemExit(f"[sf] missing file / 文件不存在: {p}")
    except json.JSONDecodeError as exc:
        raise SystemExit(f"[sf] invalid JSON in {p}: {exc} / JSON 格式错误")


def save_json(path: Path | str, data: Any) -> None:
    """
    Write an object as pretty-printed UTF-8 JSON with Chinese kept readable
    (ensure_ascii=False), a trailing newline, and parent directories created.
    All workspace JSON files go through this so git diffs stay stable.

    以缩进、保留中文可读（ensure_ascii=False）、末尾换行的方式写出 JSON，并自动创建父
    目录。workspace 中所有 JSON 都通过此函数写出，保证 git diff 稳定。
    """
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
        fh.write("\n")


def read_text(path: Path | str) -> str:
    """
    Read a UTF-8 text file, exiting with a clear message when it is missing.

    读取 UTF-8 文本文件，文件不存在时给出清晰提示并退出。
    """
    p = Path(path)
    if not p.exists():
        raise SystemExit(f"[sf] missing file / 文件不存在: {p}")
    return p.read_text(encoding="utf-8")


def write_text(path: Path | str, text: str) -> None:
    """
    Write UTF-8 text, creating parent directories, guaranteeing one trailing
    newline.

    写出 UTF-8 文本，自动创建父目录，并保证末尾恰好一个换行。
    """
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(text.rstrip("\n") + "\n", encoding="utf-8")


def sha256_of(path: Path | str) -> str | None:
    """
    Return the SHA-256 hex digest of a file, or None when the file does not
    exist. Checkpoints in state.json store these so a later `status` call can
    tell which outputs the user edited during a stop.

    返回文件的 SHA-256 十六进制摘要；文件不存在时返回 None。state.json 的检查点保存这些
    摘要，之后 `status` 命令据此判断用户在停顿期间修改了哪些产出文件。
    """
    p = Path(path)
    if not p.exists():
        return None
    h = hashlib.sha256()
    with p.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def content_hash(path: Path | str) -> str | None:
    """
    Return a hash of a file's *content* for checkpoints: for .xlsx files the
    SHA-256 of every sheet's cell values (None normalised to ""), for every
    other file the byte SHA-256. Re-saving a workbook in Excel changes its
    bytes but not its cells, so checkpoints based on this hash only flag real
    edits.

    返回用于检查点的文件*内容*哈希：.xlsx 取各表全部单元格值（None 归一为 ""）的 SHA-256，
    其他文件取字节 SHA-256。用 Excel 打开再保存会改变字节但不改变单元格，因此基于该哈希的
    检查点只会提示真正的修改。
    """
    p = Path(path)
    if not p.exists():
        return None
    if p.suffix.lower() != ".xlsx":
        return sha256_of(p)
    from openpyxl import load_workbook  # local import: keep sf_common light for non-xlsx use
    h = hashlib.sha256()
    wb = load_workbook(p, data_only=True, read_only=True)
    for name in wb.sheetnames:
        h.update(name.encode("utf-8"))
        for row in wb[name].iter_rows(values_only=True):
            h.update(json.dumps(["" if v is None else v for v in row], ensure_ascii=False, default=str).encode("utf-8"))
    return h.hexdigest()


def say(en: str, zh: str = "", *, prefix: str = "[sf]") -> None:
    """
    Print one bilingual console line: English first (what the agent reads),
    then the Chinese rendering after a slash when provided.

    打印一行双语控制台信息：英文在前（供 agent 阅读），若提供中文则以斜杠分隔跟在后面。
    """
    line = f"{prefix} {en}" if en else prefix
    if zh:
        line += f"  /  {zh}"
    print(line)


def fail(en: str, zh: str = "", code: int = 1) -> None:
    """
    Print a bilingual error line to stderr and exit with the given code.

    向 stderr 打印双语错误信息并以指定退出码退出。
    """
    msg = f"[sf:error] {en}"
    if zh:
        msg += f"  /  {zh}"
    print(msg, file=sys.stderr)
    raise SystemExit(code)


def stage_index(stage: str) -> int:
    """
    Return the 0-based position of a stage name in STAGES, failing loudly for
    unknown names such as "S9" or "s3".

    返回阶段名在 STAGES 中的下标（从 0 开始），遇到 "S9"、"s3" 之类的非法名称时直接报错。
    """
    if stage not in STAGES:
        fail(f"unknown stage {stage!r}; expected one of {STAGES}", f"未知阶段 {stage!r}")
    return STAGES.index(stage)


def resolve_workspace(arg: str | None) -> Path:
    """
    Resolve the workspace directory used by a script: an explicit --workspace
    argument wins; otherwise the SF_WORKSPACE environment variable; otherwise
    the name stored in workspace/ACTIVE by `sf_state.py init`. A bare name is
    interpreted relative to workspace/. This is what makes one-command resume
    possible after a session dies.

    解析脚本使用的 workspace 目录：优先使用显式的 --workspace 参数，其次是环境变量
    SF_WORKSPACE，最后读取 `sf_state.py init` 写入 workspace/ACTIVE 的名字。只给名字时
    相对于 workspace/ 解析。这是会话中断后能"一键恢复"的基础。
    """
    candidate = arg or os.environ.get("SF_WORKSPACE")
    if not candidate and ACTIVE_FILE.exists():
        candidate = ACTIVE_FILE.read_text(encoding="utf-8").strip()
    if not candidate:
        fail("no workspace given and workspace/ACTIVE missing; run sf_state.py init first",
             "未指定 workspace 且 workspace/ACTIVE 不存在，请先运行 sf_state.py init")
    p = Path(candidate)
    if not p.is_absolute() and not p.exists():
        p = WORKSPACE_ROOT / candidate
    return p.resolve()


def next_id(existing: Iterable[str], prefix: str, width: int = 3) -> str:
    """
    Produce the next identifier after the highest numbered one already used,
    e.g. next_id(["T001","T007"], "T") -> "T008". Ids are never reused, so
    dropped terms and closed questions keep their numbers forever.

    根据已使用的最大编号生成下一个 id，例如 next_id(["T001","T007"], "T") -> "T008"。
    id 永不复用，因此被丢弃的术语、已关闭的问题会永久保留其编号。
    """
    pat = re.compile(rf"^{re.escape(prefix)}(\d+)$")
    highest = 0
    for item in existing:
        m = pat.match(str(item))
        if m:
            highest = max(highest, int(m.group(1)))
    return f"{prefix}{highest + 1:0{width}d}"


def split_list(cell: Any) -> list[str]:
    """
    Turn an Excel cell holding "a; b; c" (or None) into ["a","b","c"], used
    when reading list-valued columns back from xlsx.

    把 Excel 单元格中的 "a; b; c"（或空值）转换成 ["a","b","c"]，用于从 xlsx 读回列表列。
    """
    if cell is None:
        return []
    if isinstance(cell, list):
        return [str(x).strip() for x in cell if str(x).strip()]
    return [s.strip() for s in re.split(r"[;；]\s*", str(cell)) if s.strip()]


def join_list(values: Iterable[str] | None) -> str:
    """
    Inverse of split_list: render a list as "a; b; c" for an Excel cell.

    split_list 的逆操作：把列表渲染成 "a; b; c" 以写入 Excel 单元格。
    """
    return LIST_SEP.join(str(v) for v in (values or []))


def relpath(p: Path | str) -> str:
    """
    Render a path relative to the repo root for compact console output.

    以相对仓库根目录的形式渲染路径，使控制台输出更紧凑。
    """
    try:
        return str(Path(p).resolve().relative_to(REPO_ROOT))
    except ValueError:
        return str(p)
